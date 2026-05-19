"""Excel (.xlsx / .xlsm) content export. One Markdown table per sheet; formulas and values shown side-by-side."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook  # type: ignore

from .ooxml_pack import OOXMLPackage
from .media import extract_media
from .embedded import extract_embedded
from .metadata import extract_metadata
from .charts import extract_charts


MAX_ROWS = 500
MAX_COLS = 50


_REF_RE = re.compile(r"^\s*(?:'([^']+)'|([A-Za-z_][\w]*))!(\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)\s*$")


def _make_resolver(wb_v: Any):
    """Callback turning a chart c:f reference into a list of cached cell values."""
    def resolver(formula: str) -> list[str]:
        m = _REF_RE.match(formula)
        if not m:
            return []
        sheet_name = m.group(1) or m.group(2)
        rng = m.group(3).replace("$", "")
        if sheet_name not in wb_v.sheetnames:
            return []
        ws = wb_v[sheet_name]
        try:
            cells = ws[rng]
        except Exception:
            return []
        out: list[str] = []
        # Flatten single cell / 1D / 2D into a flat list of strings.
        if hasattr(cells, "value") and not isinstance(cells, tuple):
            out.append("" if cells.value is None else str(cells.value))
            return out
        for item in cells:
            if isinstance(item, tuple):
                for c in item:
                    out.append("" if c.value is None else str(c.value))
            else:
                out.append("" if item.value is None else str(item.value))
        return out
    return resolver


def _cell_text(formula_cell: Any, value_cell: Any) -> str:
    fv = formula_cell.value
    vv = value_cell.value
    if isinstance(fv, str) and fv.startswith("="):
        if vv is None:
            return fv.replace("|", "\\|")
        return f"{fv} -> {vv}".replace("|", "\\|")
    if vv is None:
        return ""
    s = str(vv)
    return s.replace("|", "\\|").replace("\n", " ")


def _sheet_md(ws_formula: Any, ws_value: Any) -> tuple[str, dict[str, Any]]:
    max_row = ws_formula.max_row or 0
    max_col = ws_formula.max_column or 0
    info = {
        "name": ws_formula.title,
        "rows": max_row,
        "cols": max_col,
        "truncated": False,
    }
    if max_row == 0 or max_col == 0:
        return f"## Sheet: {ws_formula.title}\n\n_(empty)_\n", info

    use_rows = min(max_row, MAX_ROWS)
    use_cols = min(max_col, MAX_COLS)
    info["truncated"] = (use_rows < max_row) or (use_cols < max_col)

    grid: list[list[str]] = []
    for r in range(1, use_rows + 1):
        row = []
        for c in range(1, use_cols + 1):
            row.append(_cell_text(ws_formula.cell(r, c), ws_value.cell(r, c)))
        grid.append(row)

    # Use first row as Markdown table header.
    header = grid[0] if grid else []
    lines = [f"## Sheet: {ws_formula.title}", ""]
    if info["truncated"]:
        lines.append(f"_(truncated: original {max_row} rows x {max_col} cols; showing first {use_rows} x {use_cols})_")
        lines.append("")
    lines.append("| " + " | ".join(header) + " |")
    lines.append("| " + " | ".join(["---"] * len(header)) + " |")
    for row in grid[1:]:
        lines.append("| " + " | ".join(row) + " |")
    lines.append("")
    return "\n".join(lines), info


def dump(input_path: Path, out_dir: Path, options: dict[str, Any]) -> dict[str, Any]:
    out_dir.mkdir(parents=True, exist_ok=True)
    manifest: dict[str, Any] = {
        "doctype": "xlsx",
        "source": input_path.name,
        "warnings": [],
    }

    try:
        wb_f = load_workbook(str(input_path), data_only=False)
        wb_v = load_workbook(str(input_path), data_only=True)
    except Exception as e:
        manifest["warnings"].append(f"openpyxl open failed: {e}")
        wb_f = wb_v = None

    with OOXMLPackage(input_path) as pkg:
        manifest["metadata"] = extract_metadata(pkg)
        manifest["media"] = extract_media(pkg, out_dir / "media", "xlsx")
        manifest["embedded"] = extract_embedded(pkg, out_dir / "embedded", "xlsx")
        if not options.get("no_charts"):
            resolver = _make_resolver(wb_v) if wb_v is not None else None
            chart_items, chart_warns = extract_charts(pkg, out_dir / "charts", "xlsx", resolver=resolver)
            manifest["charts"] = chart_items
            manifest["warnings"].extend(chart_warns)

    md_lines = [f"# {input_path.stem}", ""]
    sheet_infos: list[dict[str, Any]] = []
    if wb_f is not None and wb_v is not None:
        # defined names
        try:
            names = []
            for n in wb_f.defined_names:
                dn = wb_f.defined_names[n]
                names.append({"name": n, "value": str(dn.value) if dn.value is not None else ""})
            if names:
                manifest["defined_names"] = names
                md_lines.append("## Defined names")
                for nm in names:
                    md_lines.append(f"- `{nm['name']}` = `{nm['value']}`")
                md_lines.append("")
        except Exception as e:
            manifest["warnings"].append(f"defined_names parse failed: {e}")

        for sheet_name in wb_f.sheetnames:
            try:
                ws_f = wb_f[sheet_name]
                ws_v = wb_v[sheet_name]
                md, info = _sheet_md(ws_f, ws_v)
                md_lines.append(md)
                sheet_infos.append(info)
            except Exception as e:
                manifest["warnings"].append(f"sheet {sheet_name} failed: {e}")
        try:
            wb_f.close()
            wb_v.close()
        except Exception:
            pass
    else:
        md_lines.append("_(openpyxl could not open the file; body not extracted)_")

    manifest["sheets"] = sheet_infos
    (out_dir / "content.md").write_text("\n".join(md_lines), encoding="utf-8")
    return manifest
