"""PowerPoint (.pptx / .pptm) content export plus rendered per-slide PNGs."""
from __future__ import annotations

import json
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

from pptx import Presentation  # type: ignore

from .ooxml_pack import OOXMLPackage
from .media import extract_media
from .embedded import extract_embedded
from .metadata import extract_metadata
from .charts import extract_charts
from .pptx_renderer import render_slides


P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _shape_text(shape: Any) -> str:
    if not shape.has_text_frame:
        return ""
    tf = shape.text_frame
    out: list[str] = []
    for para in tf.paragraphs:
        text = "".join(run.text or "" for run in para.runs)
        if text.strip():
            out.append(text)
    return "\n".join(out)


def _table_md(shape: Any) -> str:
    try:
        tbl = shape.table
    except Exception:
        return ""
    rows: list[list[str]] = []
    for row in tbl.rows:
        cells = []
        for cell in row.cells:
            t = (cell.text or "").replace("|", "\\|").replace("\n", " <br> ")
            cells.append(t)
        rows.append(cells)
    if not rows:
        return ""
    col_count = max(len(r) for r in rows)
    rows = [r + [""] * (col_count - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |",
             "| " + " | ".join(["---"] * col_count) + " |"]
    for r in rows[1:]:
        lines.append("| " + " | ".join(r) + " |")
    return "\n".join(lines)


def _slide_md(slide: Any, idx: int) -> str:
    lines = [f"## Slide {idx}", ""]

    # Title placeholder, if any.
    title = None
    title_id = None
    try:
        title = slide.shapes.title
        if title is not None and (title.text or "").strip():
            lines.append(f"### {title.text.strip()}")
            lines.append("")
            try:
                title_id = title.shape_id
            except Exception:
                title_id = id(title)
    except Exception:
        title = None

    # Walk shapes in z-order (python-pptx already yields them in z-order).
    for shape in slide.shapes:
        try:
            sid = shape.shape_id
        except Exception:
            sid = id(shape)
        if title_id is not None and sid == title_id:
            continue
        if shape.has_table:
            tbl_md = _table_md(shape)
            if tbl_md:
                lines.append(tbl_md)
                lines.append("")
            continue
        text = _shape_text(shape)
        if text.strip():
            lines.append(text)
            lines.append("")

    # Speaker notes.
    try:
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes and notes.strip():
                lines.append("**Notes**:")
                lines.append(notes.strip())
                lines.append("")
    except Exception:
        pass

    return "\n".join(lines)


def _extract_pptx_comments(pkg: OOXMLPackage) -> list[dict[str, Any]]:
    """ppt/comments/comment*.xml + ppt/commentAuthors.xml"""
    authors: dict[str, str] = {}
    if pkg.has("ppt/commentAuthors.xml"):
        try:
            root = ET.fromstring(pkg.read("ppt/commentAuthors.xml"))
            for a in root.iter(f"{P}cmAuthor"):
                authors[a.attrib.get("id", "")] = a.attrib.get("name", "")
        except Exception:
            pass

    items: list[dict[str, Any]] = []
    for name in sorted(pkg.names()):
        if not (name.startswith("ppt/comments/") and name.endswith(".xml")):
            continue
        try:
            root = ET.fromstring(pkg.read(name))
        except Exception:
            continue
        for c in root.iter(f"{P}cm"):
            text_el = c.find(f"{P}text")
            items.append({
                "src": name,
                "author_id": c.attrib.get("authorId", ""),
                "author": authors.get(c.attrib.get("authorId", ""), ""),
                "date": c.attrib.get("dt", ""),
                "text": (text_el.text or "").strip() if text_el is not None else "",
            })
    return items


def dump(input_path: Path, out_dir: Path, options: dict[str, Any]) -> dict[str, Any]:
    out_dir.mkdir(parents=True, exist_ok=True)
    manifest: dict[str, Any] = {
        "doctype": "pptx",
        "source": input_path.name,
        "warnings": [],
    }

    try:
        pres = Presentation(str(input_path))
    except Exception as e:
        manifest["warnings"].append(f"python-pptx open failed: {e}")
        pres = None

    with OOXMLPackage(input_path) as pkg:
        manifest["metadata"] = extract_metadata(pkg)
        manifest["media"] = extract_media(pkg, out_dir / "media", "pptx")
        manifest["embedded"] = extract_embedded(pkg, out_dir / "embedded", "pptx")
        if not options.get("no_charts"):
            chart_items, chart_warns = extract_charts(pkg, out_dir / "charts", "pptx")
            manifest["charts"] = chart_items
            manifest["warnings"].extend(chart_warns)

        comments = _extract_pptx_comments(pkg)

    md_lines = [f"# {input_path.stem}", ""]
    slide_count = 0
    if pres is not None:
        for i, slide in enumerate(pres.slides, start=1):
            slide_count += 1
            try:
                md_lines.append(_slide_md(slide, i))
            except Exception as e:
                manifest["warnings"].append(f"slide {i} parse failed: {e}")
    manifest["slides_count"] = slide_count

    (out_dir / "content.md").write_text("\n".join(md_lines), encoding="utf-8")

    if comments:
        (out_dir / "comments.json").write_text(
            json.dumps(comments, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        manifest["comments"] = {"count": len(comments), "file": "comments.json"}

    # Render slides.
    renderer_pref = options.get("slide_renderer", "auto")
    if renderer_pref != "none":
        render_result = render_slides(
            input_path, out_dir,
            dpi=int(options.get("dpi", 150)),
            prefer=renderer_pref,
        )
        manifest["render"] = {
            "renderer": render_result["renderer"],
            "slides": render_result["slides"],
            "pdf": render_result["pdf"],
        }
        manifest["warnings"].extend(render_result["warnings"])

    return manifest
